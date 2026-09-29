"""Rebuild the final report using the Codex bundled Python and matplotlib.
Run from the repository root. Raw model evidence remains in its original location.
"""
import sys, json, csv, shutil, re, sqlite3, ast, hashlib
from pathlib import Path
ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / 'tmp/doc_tools/python_libs'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
A = Path(__file__).parent
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'savefig.facecolor':'white'})
def load(p): return json.loads((ROOT/p).read_text(encoding='utf-8'))
transfer=load('python_models/yamnet_transfer_metrics.json'); aug=load('python_models/yamnet_augmented_metrics.json'); smoke=load('reports/delivery_e2e.json')
classes=transfer['test']['classes']; labels=[c.replace('_',' ').replace('person asking for help','Call for help').capitalize() for c in classes]
def save(fig,name):
    fig.savefig(A/(name+'.png'),dpi=190,bbox_inches='tight'); fig.savefig(A/(name+'.svg'),bbox_inches='tight'); plt.close(fig)
def matrix(m,name,title):
    a=np.array(m['confusion_matrix']); fig,ax=plt.subplots(figsize=(8,6.3)); ax.imshow(a,cmap='Blues',vmin=0,vmax=max(a.max(),1))
    for i in range(10):
        for j in range(10): ax.text(j,i,str(a[i,j]),ha='center',va='center',fontsize=9,color='white' if a[i,j]>a.max()*.55 else '#202020')
    ax.set(xticks=range(10),yticks=range(10),xticklabels=labels,yticklabels=labels,xlabel='Predicted class',ylabel='True class',title=title)
    plt.setp(ax.get_xticklabels(),rotation=50,ha='right'); save(fig,name)
matrix(transfer['test'],'confusion_python','Active Python model | 471 test segments')
matrix(aug['test'],'confusion_augmented','Augmented candidate | 471 test segments')
gtm_cm=np.zeros((10,10),int)
for row in smoke['model_comparison']: gtm_cm[classes.index(row['truth']),classes.index(row['gtm'])]+=1
matrix({'confusion_matrix':gtm_cm},'confusion_gtm_smoke','GTM integration smoke test | only 10 recordings')
fig,ax=plt.subplots(figsize=(8,4)); y=np.arange(10)
ax.barh(y-.18,[transfer['test']['per_class'][c]['recall']*100 for c in classes],.35,label='Active Python',color='#235d7a')
ax.barh(y+.18,[aug['test']['per_class'][c]['recall']*100 for c in classes],.35,label='Augmented candidate',color='#c28636')
ax.axvline(85,color='#6d6d6d',ls='--',label='Critical recall target 85%'); ax.set(yticks=y,yticklabels=labels,xlim=(0,105),xlabel='Recall (%)'); ax.invert_yaxis(); ax.legend(loc='lower left',bbox_to_anchor=(0,1),ncol=2,frameon=False); save(fig,'class_recall')
names=['Random forest','Extra trees','Gradient boosting','Handcrafted SVM','YAMNet plus SVM']; vals=[]
for f in ['random_forest','extra_trees','gradient_boosting','svm_c10_gscale','yamnet_transfer']: vals.append(load('python_models/'+f+'_metrics.json')['validation']['accuracy']*100)
fig,ax=plt.subplots(figsize=(8,3.4)); bars=ax.barh(names,vals,color=['#9ab2bf']*4+['#235d7a']); ax.invert_yaxis();ax.set(xlim=(0,102),xlabel='Validation accuracy (%) | 477 segments');ax.bar_label(bars,fmt='%.2f%%',padding=4);save(fig,'validation_models')
fig,ax=plt.subplots(figsize=(8,3.2)); bars=ax.bar(['Clean','20 dB Gaussian noise','10 dB Gaussian noise'],[9,4,1],color=['#235d7a','#c28636','#a75048']);ax.bar_label(bars,labels=['9/10','4/10','1/10'],padding=4);ax.set(ylim=(0,11),ylabel='Correct predictions out of 10',title='Small diagnostic probe of the active Python model');save(fig,'noise_probe')
fig,ax=plt.subplots(figsize=(8,3.5)); x=np.arange(10); ax.bar(x-.18,[r['python_confidence']*100 for r in smoke['model_comparison']],.35,label='Python',color='#235d7a');ax.bar(x+.18,[r['gtm_confidence']*100 for r in smoke['model_comparison']],.35,label='GTM export',color='#c28636');ax.set(xticks=x,xticklabels=labels,ylim=(0,110),ylabel='Top class confidence (%)');plt.setp(ax.get_xticklabels(),rotation=40,ha='right');ax.legend(frameon=False,ncol=2);save(fig,'confidence_smoke')
def diagram(name,nodes,edges,size=(9,4.6)):
    fig,ax=plt.subplots(figsize=size);ax.set(xlim=(0,10),ylim=(0,6));ax.axis('off')
    for key,(x,y,w,h,label) in nodes.items():
        ax.add_patch(FancyBboxPatch((x-w/2,y-h/2),w,h,boxstyle='round,pad=0.05',fc='#eef4f6',ec='#658493',lw=1.2));ax.text(x,y,label,ha='center',va='center',fontsize=10)
    for a,b,label in edges:
        x,y,w,h,_=nodes[a];xx,yy,ww,hh,_=nodes[b];dx=xx-x;dy=yy-y
        if name=='data_flow' and a=='d' and b=='r':
            ax.plot([x,x,xx],[y-h/2,.15,.15],color='#50616d',lw=1.2)
            ax.annotate('',xy=(xx,yy-hh/2),xytext=(xx,.15),arrowprops=dict(arrowstyle='->',color='#50616d',lw=1.2))
            ax.text(5,.22,label,ha='center',fontsize=8,bbox=dict(fc='white',ec='none',pad=.5));continue
        t=min(w/(2*abs(dx)) if dx else 1e9,h/(2*abs(dy)) if dy else 1e9)
        tt=min(ww/(2*abs(dx)) if dx else 1e9,hh/(2*abs(dy)) if dy else 1e9)
        start=(x+t*dx,y+t*dy);end=(xx-tt*dx,yy-tt*dy)
        ax.annotate('',xy=end,xytext=start,arrowprops=dict(arrowstyle='->',color='#50616d',lw=1.2));
        if label:ax.text((start[0]+end[0])/2+.08,(start[1]+end[1])/2+.14,label,ha='center',fontsize=8,bbox=dict(fc='white',ec='none',pad=.5))
    save(fig,name)
diagram('architecture',{'u':(1.4,5,2.3,.8,'Browser\nUpload or microphone'),'w':(5,5,2.7,.8,'Flask application\nAuthentication and routes'),'p':(2,3,3,.85,'Python classifier\nYAMNet plus SVM'),'g':(7.8,3,3,.85,'Independent GTM export\nLog mel CNN'),'r':(5,1.5,2.7,.8,'Comparison and rules\nReview or alert'),'d':(1.6,.4,2.8,.7,'SQLite and audio files'),'v':(8.2,.4,2.8,.7,'Dashboard and reports')},[('u','w','request'),('w','p','audio'),('w','g','audio'),('p','r','scores'),('g','r','scores'),('r','d','store'),('r','v','display')])
diagram('data_flow',{'u':(1.1,5,1.9,.8,'User'),'p1':(4,5,2.6,.8,'1 Validate audio'),'a':(8,5,2.8,.8,'D1 Protected audio'),'p2':(4,3,2.6,.8,'2 Process and infer'),'m':(8,3,2.8,.8,'D2 Model artifacts'),'p3':(4,1,2.6,.8,'3 Compare and decide'),'d':(8,1,2.8,.8,'D3 Event database'),'r':(1.1,1,1.9,.8,'Reviewer')},[('u','p1','audio'),('p1','a','file'),('p1','p2','samples'),('m','p2','weights'),('p2','p3','scores'),('p3','d','event'),('d','r','review queue'),('r','p3','decision')])
diagram('use_cases',{'u':(1.3,5,2.2,.7,'Normal user'),'r':(1.3,3,2.2,.7,'Reviewer / operator'),'m':(1.3,1,2.2,.7,'Maintenance'),'a':(8.7,3,2.2,.7,'Administrator'),'c1':(5,5,3,.8,'Account, upload, monitor\nOwn history and reports'),'c2':(5,3,3,.8,'Listen and review\nCorrect class and comment'),'c3':(5,1,3,.8,'Acknowledge alerts\nInspect event history'),'c4':(8.7,5,2.3,.8,'Rules, CSV export\nRetention settings')},[('u','c1',''),('r','c2',''),('r','c3',''),('m','c3',''),('a','c2',''),('a','c3',''),('a','c4','')])
diagram('activity',{'a':(1.5,5,2.4,.7,'Sign in and choose audio'),'b':(5,5,2.7,.7,'Validate and check quality'),'e':(8.5,5,2.2,.7,'Reject invalid input'),'c':(5,3,2.7,.8,'Prepare each segment\nRun both classifiers'),'d':(5,1,2.7,.8,'Store results and rules\nShow analysis'),'r':(1.5,1,2.4,.8,'Listen and review\nRecord correction'),'f':(8.5,1,2.2,.8,'Acknowledge alert\nRecord action')},[('a','b',''),('b','e','invalid'),('b','c','usable'),('c','d',''),('d','r','uncertain'),('d','f','confirmed')])
diagram('decision',{'s':(2,5,3,.7,'Select primary model'),'q':(2,3.6,3,.8,'Both available and agree?\nQuality and scores eligible?'),'r':(7.4,3.6,3.8,.8,'Flag manual review\nKeep both original predictions'),'e':(2,2,3,.8,'Critical category and\nconsecutive count reached?'),'c':(7.4,2,3.8,.8,'Store classified event\nContinue monitoring'),'a':(2,.5,3,.7,'Generate visible alert')},[('s','q',''),('q','r','No'),('q','e','Yes'),('e','c','No'),('e','a','Yes')])
diagram('entity_relationships',{'u':(1.7,4.6,2.9,1.2,'users\nPK user_id'),'a':(6,4.6,3,1.2,'audio_files\nPK audio_id\nFK uploaded_by'),'d':(6,1.4,3,1.3,'detections\nPK detection_id\nFK audio_id, reviewed_by'),'l':(1.7,1.4,2.9,1.2,'audit_log\nPK audit_id\nFK user_id')},[('u','a','1 to many uploads'),('a','d','1 to many segments'),('u','l','1 to many actions'),('u','d','1 to many reviews')])
fig,ax=plt.subplots(figsize=(9,4.7));ax.set(xlim=(0,10),ylim=(0,10));ax.axis('off');xs=[.9,3.6,6.4,9.1]
for x,s in zip(xs,['Browser','Flask service','Models and rules','SQLite / files']):ax.text(x,9.5,s,ha='center',weight='bold');ax.plot([x,x],[.5,9],ls='--',color='#9caab1')
for y,a,b,s in [(8.3,0,1,'POST audio with session and CSRF'),(7.2,1,3,'Validate and save audio'),(6.1,1,2,'Prepare segment and infer independently'),(5,2,1,'Two score vectors and rule outcome'),(3.9,1,3,'Insert detection and audit action'),(2.8,1,0,'Return analysis page'),(1.7,0,1,'Authorized review or acknowledgement'),(.7,1,3,'Preserve scores and record human action')]:
    ax.annotate('',xy=(xs[b],y),xytext=(xs[a],y),arrowprops=dict(arrowstyle='->',color='#235d7a'));ax.text((xs[a]+xs[b])/2,y+.16,s,ha='center',fontsize=9)
save(fig,'sequence')
# Export the values behind the charts for review and reuse.
with (A/'model_results.csv').open('w',newline='',encoding='utf-8') as f:
    w=csv.writer(f);w.writerow(['model','split','accuracy','macro_precision','macro_recall','macro_f1','segments'])
    for name,m in [('active_python',transfer),('augmented_candidate',aug)]:
        for s in ['validation','test']:
            z=m[s];w.writerow([name,s,z['accuracy'],z['precision'],z['recall'],z['macro_f1'],m['split_counts'][s]])
with (A/'class_results.csv').open('w',newline='',encoding='utf-8') as f:
    w=csv.writer(f);w.writerow(['class','precision','recall','f1','support','false_positive','false_negative'])
    cm=np.array(transfer['test']['confusion_matrix'])
    for i,c in enumerate(classes):
        z=transfer['test']['per_class'][c];w.writerow([c,z['precision'],z['recall'],z['f1-score'],int(z['support']),int(cm[:,i].sum()-cm[i,i]),int(cm[i,:].sum()-cm[i,i])])
originals=load('audio_dataset/yamnet_augmented/originals_manifest.json')
counts={c:sum(r['class_label']==c for r in originals) for c in classes}
(A/'dataset_counts.json').write_text(json.dumps(counts,indent=2))
required_dirs='src templates static data audio_dataset notebooks python_models gtm_model feature_extraction audio_preprocessing augmentation alert_rules database tests sample_audio documentation screenshots reports config'.split()
for folder in ['sample_audio','screenshots']:
    p=ROOT/folder/'README.md'
    if not p.exists(): p.write_text('# '+folder.replace('_',' ').title()+'\n\n'+('Place licensed, non-confidential evaluator audio here with a source and license manifest. No audio is supplied by this documentation update; dataset rights must be verified before redistribution.' if folder=='sample_audio' else 'Place verified application screenshots here for the submission and demonstration. Report diagrams and charts are stored in documentation/assets/final_report/. Diagrams are not evidence of completed browser tests.')+'\n',encoding='utf-8')
audit={'checked_on':'2026-09-29','required_files':{n:(ROOT/n).is_file() for n in ['README.md','AI_USAGE.md','requirements.txt','LICENSE']},'required_directories':{n:(ROOT/n).is_dir() for n in required_dirs},'note':'Directory presence does not establish completeness of datasets, screenshots, video, training evidence or unseen evaluation.'}
(A/'structure_audit.json').write_text(json.dumps(audit,indent=2))
# Style and assembly.
doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72);sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6);doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Title'].font.size=Pt(30);doc.styles['Heading 1'].font.size=Pt(19);doc.styles['Heading 2'].font.size=Pt(13);doc.styles['Caption'].font.size=Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(12)
for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
footer=sec.footer.paragraphs[0];footer.alignment=2;r=footer.add_run('SonicSentinel AI  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def table(rows):
    t=doc.add_table(rows=1,cols=len(rows[0]));t.autofit=False
    widths=([2,4.83] if len(rows[0])==2 else [2,1.1,3.73] if len(rows[0])==3 else [1.85]+[(6.83-1.85)/(len(rows[0])-1)]*(len(rows[0])-1))
    if rows[0][0]=='SRS target': widths=[2.15,2.2,2.48]
    if rows[0][0]=='SRS reference': widths=[2.15,1.65,3.03]
    if rows[0][0]=='Metric': widths=[2.83,2,2]
    if rows[0][0]=='Class' and len(rows[0])==3: widths=[3.23,1.8,1.8]
    if len(rows[0])==3 and rows[0][1]=='Type and constraints': widths=[2.25,1.3,3.28]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for c,v,w in zip(cells,row,widths):
            c.width=Inches(w);c.text=str(v);pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                n=OxmlElement('w:'+side);n.set(qn('w:w'),'80');n.set(qn('w:type'),'dxa');mar.append(n)
            pr.append(mar)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.02
                for r in p.runs:r.font.size=Pt(9);r.bold=i==0
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF2');pr.append(shade)
        trPr=t.rows[i]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if i==0:trPr.append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders);doc.add_paragraph().paragraph_format.space_after=Pt(0)
text=(A/'report_content.md').read_text(encoding='utf-8')
assert '\u2014' not in text
text=text.replace('{{CLASS_TABLE}}','\n'.join('| '+ ' | '.join(map(str,r))+' |' for r in [['Class','Precision','Recall','F1','Test segments']]+[[labels[i],f"{transfer['test']['per_class'][c]['precision']:.3f}",f"{transfer['test']['per_class'][c]['recall']:.3f}",f"{transfer['test']['per_class'][c]['f1-score']:.3f}",int(transfer['test']['per_class'][c]['support'])] for i,c in enumerate(classes)]))
text=text.replace('{{DATA_TABLE}}','\n'.join('| '+' | '.join(map(str,r))+' |' for r in [['Source class','Originals']]+[[labels[i],counts[c]] for i,c in enumerate(classes)]))
text=text.replace('{{ERROR_TABLE}}','\n'.join('| '+' | '.join(map(str,r))+' |' for r in [['Class','False positives','False negatives']]+[[labels[i],int(cm[:,i].sum()-cm[i,i]),int(cm[i,:].sum()-cm[i,i])] for i in range(10)]))
lines=text.splitlines();i=0;fig_no=0
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line=='---PAGE---':doc.add_page_break();continue
    if line.startswith('|'):
        rows=[[v.strip() for v in line.strip('|').split('|')]]
        while i<len(lines) and lines[i].strip().startswith('|'):rows.append([v.strip() for v in lines[i].strip().strip('|').split('|')]);i+=1
        table(rows);continue
    if line.startswith('!['):
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line);fig_no+=1;p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
        width=6.0 if m[2] in ['confidence_smoke.png','noise_probe.png'] else 6.6
        pic=p.add_run().add_picture(str(A/m[2]),width=Inches(width));pic._inline.docPr.set('descr',m[1]);doc.add_paragraph(f'Figure {fig_no}. '+m[1],style='Caption');continue
    if line.startswith('# '):doc.add_paragraph(line[2:],style='Title');continue
    if line.startswith('## '):doc.add_heading(line[3:],1);continue
    if line.startswith('### '):doc.add_heading(line[4:],2);continue
    if line.startswith('- '):doc.add_paragraph(line[2:],style='List Bullet');continue
    if line.startswith('> '):
        p=doc.add_paragraph();r=p.add_run(line[2:]);r.font.name='Consolas';r.font.size=Pt(9);continue
    doc.add_paragraph(line)
doc.core_properties.title='SonicSentinel AI Project Documentation';doc.core_properties.subject='Architecture, models, evaluation, operation and SRS assessment';doc.core_properties.author='SonicSentinel AI project team';doc.core_properties.keywords='audio classification, project documentation, SRS, evaluation'
out=ROOT/'documentation/documentation.docx';doc.save(out)
draft=ROOT/'documentation/SonicSentinelAI Documentation.docx';backup=A/'original_draft.docx'
if not backup.exists():shutil.copy2(draft,backup)
shutil.copy2(out,draft)
(A/'README.md').write_text('# Final report assets\n\nGenerated 2026-09-29. `report_content.md` and `build_report.py` reproduce the Word report. PNG and SVG files are original diagrams and plots derived from the repository evidence. CSV/JSON files preserve the chart values and structure audit. `original_draft.docx` preserves the supplied draft. QA render files are stored separately under `tmp/`.\n\nModel charts use stored evaluation outputs, not newly trained models. The GTM confusion matrix contains only ten integration examples. No training curves are fabricated.\n',encoding='utf-8')
print(out);print('Figures:',fig_no,'paragraphs:',len(doc.paragraphs),'tables:',len(doc.tables))
